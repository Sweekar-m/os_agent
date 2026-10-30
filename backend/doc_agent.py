"""
Universal Document Agent
Supports: PowerPoint (.pptx), Word (.docx), Excel (.xlsx), PDF (.pdf)
Features: Read, create, edit, summarize, RAG index
"""

import os
import re
import time
import string
import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
from concurrent.futures import ThreadPoolExecutor

# ── Core doc libraries ──────────────────────────────────────────────────────
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

try:
    import docx as _docx
    from docx import Document as DocxDocument
    from docx.shared import Pt as DocxPt, RGBColor as DocxRGB, Inches as DocxInches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import (Font, PatternFill, Alignment,
                                 Border, Side, GradientFill)
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import pdfplumber
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

try:
    import reportlab
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors as r_colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    import win32com.client
    import pythoncom
    HAS_WIN32COM = True
except ImportError:
    HAS_WIN32COM = False

DOC_EXTENSIONS = {
    "pptx": "PowerPoint", "ppt":  "PowerPoint",
    "docx": "Word",        "doc":  "Word",
    "xlsx": "Excel",       "xls":  "Excel",
    "pdf":  "PDF",
}

EXCLUDE_DIRS = {
    'windows', 'program files', 'program files (x86)', 'programdata',
    '$recycle.bin', 'recycle.bin', 'system volume information', 'appdata', '.git',
    'node_modules', 'venv', '.venv', '__pycache__', '.cache', 'site-packages',
    'microsoft', 'windowsapps', '$windows.~bt', '$windows.~ws', 'recovery',
    'boot', '$sysreset', 'perflogs', '.vscode', '.gemini', 'dist', 'build'
}


def _fmt_size(n: int) -> str:
    if n < 1024:      return f"{n} B"
    if n < 1024**2:   return f"{n/1024:.1f} KB"
    return f"{n/1024**2:.1f} MB"

def _fmt_time(ts: float) -> str:
    dt  = datetime.datetime.fromtimestamp(ts)
    now = datetime.datetime.now()
    d   = (now - dt).days
    if d == 0:  return f"Today {dt.strftime('%I:%M %p')}"
    if d == 1:  return "Yesterday"
    if d < 7:   return f"{d} days ago"
    return dt.strftime('%b %d, %Y')


class DocAgent:
    def __init__(self):
        self.active_file: Optional[str] = None
        self.recent_files: List[Dict[str, Any]] = []
        self._rag_cache: Dict[str, "DocRAGIndex"] = {}
        self._last_scan: float = 0.0

    def _priority_dirs(self) -> List[Path]:
        home = Path.home()
        candidates = [
            home / "Documents", home / "Desktop", home / "Downloads",
            home / "OneDrive", home / "OneDrive" / "Documents",
            Path("D:/"), Path("D:/PROJECT"), Path("D:/project_a"), Path("C:/"),
        ]
        seen, dirs = set(), []
        for c in candidates:
            try:
                if c.exists():
                    r = str(c.resolve()).lower()
                    if r not in seen:
                        seen.add(r); dirs.append(c)
            except Exception:
                pass
        return dirs

    def find_recent_docs(
        self,
        query: Optional[str] = None,
        drive: Optional[str] = None,
        file_type: Optional[str] = None,
        limit: int = 25,
        force_refresh: bool = False
    ) -> List[Dict[str, Any]]:
        now = time.time()
        if not force_refresh and not query and not drive and not file_type \
                and self.recent_files and (now - self._last_scan < 10):
            return self.recent_files[:limit]

        wanted_exts = {file_type.lower().lstrip('.')} if file_type else set(DOC_EXTENSIONS.keys())
        found, seen = [], set()
        q_lower = query.lower().strip() if query else None

        def _add(p_str: str, fname: str):
            pl = p_str.lower()
            if pl in seen: return
            ext = fname.rsplit('.', 1)[-1].lower() if '.' in fname else ''
            if fname.startswith('~$') or ext not in wanted_exts: return
            if q_lower and (q_lower not in fname.lower() and q_lower not in pl): return
            seen.add(pl)
            try:
                st = os.stat(p_str)
                if st.st_size > 0:
                    found.append({
                        "name": fname, "path": p_str, "ext": ext,
                        "doc_type": DOC_EXTENSIONS.get(ext, "Unknown"),
                        "mtime": st.st_mtime, "size_bytes": st.st_size,
                        "size_display": _fmt_size(st.st_size),
                        "modified_display": _fmt_time(st.st_mtime),
                        "drive": p_str[:2].upper(),
                    })
            except Exception: pass

        for p_dir in self._priority_dirs():
            if drive and not str(p_dir).lower().startswith(drive.lower()): continue
            try:
                base = len(p_dir.parts)
                for root, dirs, files in os.walk(p_dir, topdown=True):
                    if len(Path(root).parts) - base > 5:
                        dirs.clear(); continue
                    dirs[:] = [d for d in dirs if d.lower() not in EXCLUDE_DIRS
                                and not d.startswith('.') and not d.startswith('$')]
                    for f in files:
                        _add(os.path.join(root, f), f)
            except Exception: pass

        drives = [drive] if drive else [f"{l}:\\" for l in string.ascii_uppercase
                                         if os.path.exists(f"{l}:\\")]

        def _scan(d_root):
            t0 = time.time()
            try:
                for root, dirs, files in os.walk(d_root, topdown=True):
                    if time.time() - t0 > 3: break
                    dirs[:] = [d for d in dirs if d.lower() not in EXCLUDE_DIRS
                                and not d.startswith('.') and not d.startswith('$')]
                    rl = root.lower()
                    if any(x in rl for x in ('$recycle', 'recycle.bin', '\\appdata\\', '\\windows\\')):
                        continue
                    for f in files:
                        _add(os.path.join(root, f), f)
            except Exception: pass

        with ThreadPoolExecutor(max_workers=max(1, len(drives))) as ex:
            list(ex.map(_scan, drives))

        found.sort(key=lambda x: x["mtime"], reverse=True)
        top = found[:limit]
        if not query and not drive and not file_type:
            self.recent_files = top; self._last_scan = now
        return top

    # ── Readers ──────────────────────────────────────────────────────────

    def read_pptx(self, path: str) -> str:
        lines = []
        try:
            prs = pptx.Presentation(path)
            lines.append(f"[PowerPoint] {Path(path).name} — {len(prs.slides)} slides\n")
            for i, slide in enumerate(prs.slides, 1):
                title = ""
                if slide.shapes.title and slide.shapes.title.has_text_frame:
                    title = slide.shapes.title.text.strip()
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            t = para.text.strip()
                            if t and t != title: texts.append(t)
                    if shape.has_table:
                        for row in shape.table.rows:
                            cells = [c.text.strip() for c in row.cells if c.text.strip()]
                            if cells: texts.append(" | ".join(cells))
                try:
                    if slide.has_notes_slide:
                        n = slide.notes_slide.notes_text_frame.text.strip()
                        if n: texts.append(f"[Notes: {n}]")
                except Exception: pass
                lines.append(f"--- Slide {i}: {title or '(no title)'} ---")
                lines.extend(texts); lines.append("")
        except Exception as e:
            lines.append(f"[Error: {e}]")
        return "\n".join(lines)

    def read_docx(self, path: str) -> str:
        if not HAS_DOCX: return "[python-docx not installed]"
        lines = []
        try:
            doc = DocxDocument(path)
            lines.append(f"[Word] {Path(path).name}\n")
            for para in doc.paragraphs:
                t = para.text.strip()
                if not t: continue
                style = para.style.name if para.style else ""
                if "Heading" in style:
                    lv = style.replace("Heading", "").strip()
                    lines.append(f"\n{'#'*int(lv) if lv.isdigit() else '##'} {t}")
                else:
                    lines.append(t)
            for i, table in enumerate(doc.tables, 1):
                lines.append(f"\n[Table {i}]")
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells: lines.append(" | ".join(cells))
        except Exception as e:
            lines.append(f"[Error: {e}]")
        return "\n".join(lines)

    def read_xlsx(self, path: str) -> str:
        if not HAS_OPENPYXL: return "[openpyxl not installed]"
        lines = []
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines.append(f"[Excel] {Path(path).name} — {len(wb.sheetnames)} sheet(s)\n")
            for sname in wb.sheetnames:
                ws = wb[sname]
                lines.append(f"\n--- Sheet: {sname} ---")
                rows_read = 0
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() if c is not None else "" for c in row]
                    if any(c for c in cells):
                        lines.append(" | ".join(cells)); rows_read += 1
                    if rows_read >= 200:
                        lines.append("... (truncated at 200 rows)"); break
            wb.close()
        except Exception as e:
            lines.append(f"[Error: {e}]")
        return "\n".join(lines)

    def read_pdf(self, path: str) -> str:
        if not HAS_PDF: return "[pdfplumber not installed]"
        lines = []
        try:
            with pdfplumber.open(path) as pdf:
                lines.append(f"[PDF] {Path(path).name} — {len(pdf.pages)} pages\n")
                for i, page in enumerate(pdf.pages, 1):
                    text = (page.extract_text() or "").strip()
                    if text:
                        lines.append(f"--- Page {i} ---"); lines.append(text); lines.append("")
                    if i >= 50:
                        lines.append("... (truncated at 50 pages)"); break
        except Exception as e:
            lines.append(f"[Error: {e}]")
        return "\n".join(lines)

    def read_file(self, path: str) -> str:
        ext = Path(path).suffix.lower().lstrip('.')
        if ext in ("pptx", "ppt"):   return self.read_pptx(path)
        if ext in ("docx", "doc"):   return self.read_docx(path)
        if ext in ("xlsx", "xls"):   return self.read_xlsx(path)
        if ext == "pdf":             return self.read_pdf(path)
        try:
            return Path(path).read_text(encoding="utf-8", errors="replace")
        except Exception as e:
            return f"[Cannot read: {e}]"

    def get_rag_index(self, path: str, force_refresh: bool = False) -> "DocRAGIndex":
        norm = os.path.abspath(path)
        cached = self._rag_cache.get(norm)
        if cached and not force_refresh and (time.time() - cached.indexed_at < 60):
            return cached
        idx = DocRAGIndex(path, self)
        self._rag_cache[norm] = idx
        return idx

    def inspect_file(self, path: str) -> Dict[str, Any]:
        p = Path(path)
        if not p.exists(): raise FileNotFoundError(f"File not found: {path}")
        ext = p.suffix.lower().lstrip('.')
        st  = p.stat()
        info: Dict[str, Any] = {
            "name": p.name, "path": str(p.resolve()),
            "doc_type": DOC_EXTENSIONS.get(ext, "Unknown"), "ext": ext,
            "size_display": _fmt_size(st.st_size),
            "modified_display": _fmt_time(st.st_mtime),
        }
        if ext in ("pptx", "ppt"):
            try:
                prs = pptx.Presentation(path)
                slides = []
                for i, slide in enumerate(prs.slides, 1):
                    title = (slide.shapes.title.text.strip()
                             if slide.shapes.title and slide.shapes.title.has_text_frame
                             else f"Slide {i}")
                    texts = [p.text.strip() for shp in slide.shapes
                             if shp.has_text_frame
                             for p in shp.text_frame.paragraphs
                             if p.text.strip() and p.text.strip() != title]
                    slides.append({"slide_num": i, "title": title,
                                   "preview": " | ".join(texts[:2])})
                info["slides"] = slides; info["total_slides"] = len(prs.slides)
            except Exception as e: info["error"] = str(e)

        elif ext in ("docx", "doc") and HAS_DOCX:
            try:
                doc = DocxDocument(path)
                headings = [p.text.strip() for p in doc.paragraphs
                            if p.style and "Heading" in p.style.name and p.text.strip()]
                info["headings"] = headings[:20]
                info["paragraph_count"] = len([p for p in doc.paragraphs if p.text.strip()])
                info["table_count"] = len(doc.tables)
            except Exception as e: info["error"] = str(e)

        elif ext in ("xlsx", "xls") and HAS_OPENPYXL:
            try:
                wb = openpyxl.load_workbook(path, read_only=True)
                info["sheets"] = wb.sheetnames; info["sheet_count"] = len(wb.sheetnames)
                wb.close()
            except Exception as e: info["error"] = str(e)

        elif ext == "pdf" and HAS_PDF:
            try:
                with pdfplumber.open(path) as pdf:
                    info["page_count"] = len(pdf.pages)
            except Exception as e: info["error"] = str(e)

        self.active_file = str(p.resolve())
        return info

    # ── Creators ─────────────────────────────────────────────────────────

    def create_ppt(self, file_path: str, title: str,
                   slides_data: List[Dict[str, Any]], theme: str = "dark") -> Dict[str, Any]:
        """Create styled PowerPoint. Delegates to ppt_agent for consistency."""
        from ppt_agent import ppt_agent as _ppt
        result = _ppt.create_ppt(file_path, title, slides_data, theme=theme)
        self.active_file = str(Path(file_path).resolve())
        return result

    def create_pdf(self, file_path: str, title: str,
                   sections: List[Dict[str, Any]], theme: str = "professional") -> Dict[str, Any]:
        """Generates a styled, multi-page professional PDF document using ReportLab."""
        if not HAS_REPORTLAB:
            return {"status": "error", "message": "reportlab not installed"}

        try:
            THEMES = {
                "professional": {
                    "primary": r_colors.HexColor("#0f172a"),
                    "secondary": r_colors.HexColor("#0284c7"),
                    "body": r_colors.HexColor("#334155"),
                    "accent": r_colors.HexColor("#10b981")
                },
                "modern": {
                    "primary": r_colors.HexColor("#1e1b4b"),
                    "secondary": r_colors.HexColor("#6366f1"),
                    "body": r_colors.HexColor("#1f2937"),
                    "accent": r_colors.HexColor("#8b5cf6")
                },
                "corporate": {
                    "primary": r_colors.HexColor("#1c1917"),
                    "secondary": r_colors.HexColor("#ea580c"),
                    "body": r_colors.HexColor("#292524"),
                    "accent": r_colors.HexColor("#f59e0b")
                }
            }
            c = THEMES.get(theme, THEMES["professional"])

            Path(file_path).parent.mkdir(parents=True, exist_ok=True)
            doc = SimpleDocTemplate(
                file_path,
                pagesize=letter,
                leftMargin=54,
                rightMargin=54,
                topMargin=54,
                bottomMargin=54
            )

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'DocTitle',
                parent=styles['Title'],
                fontName='Helvetica-Bold',
                fontSize=24,
                leading=28,
                textColor=c["primary"],
                alignment=0,
                spaceAfter=6
            )
            date_style = ParagraphStyle(
                'DocDate',
                parent=styles['Normal'],
                fontName='Helvetica',
                fontSize=10,
                leading=14,
                textColor=c["secondary"],
                spaceAfter=14
            )
            h2_style = ParagraphStyle(
                'DocH2',
                parent=styles['Heading2'],
                fontName='Helvetica-Bold',
                fontSize=14,
                leading=18,
                textColor=c["secondary"],
                spaceBefore=14,
                spaceAfter=6
            )
            body_style = ParagraphStyle(
                'DocBody',
                parent=styles['Normal'],
                fontName='Helvetica',
                fontSize=10.5,
                leading=15,
                textColor=c["body"],
                spaceAfter=8
            )
            bullet_style = ParagraphStyle(
                'DocBullet',
                parent=styles['Normal'],
                fontName='Helvetica',
                fontSize=10,
                leading=14,
                textColor=c["body"],
                leftIndent=18,
                firstLineIndent=-10,
                spaceAfter=4
            )

            story = []
            story.append(Paragraph(title, title_style))
            date_str = datetime.datetime.now().strftime("%B %d, %Y")
            story.append(Paragraph(f"Generated on {date_str} • Prepared by Windows OS Assistant", date_style))
            story.append(HRFlowable(width="100%", thickness=1.5, color=c["secondary"], spaceBefore=2, spaceAfter=14))

            for sec in sections:
                heading = sec.get("heading", "")
                if heading:
                    story.append(Paragraph(heading, h2_style))
                for p in sec.get("paragraphs", []):
                    story.append(Paragraph(p, body_style))
                for b in sec.get("bullets", []):
                    story.append(Paragraph(f"• {b}", bullet_style))
                story.append(Spacer(1, 8))

            doc.build(story)
            self.active_file = str(Path(file_path).resolve())
            return {
                "status": "success",
                "action": "create_pdf",
                "file": file_path,
                "title": title,
                "section_count": len(sections),
                "theme": theme
            }
        except Exception as e:
            return {"status": "error", "message": f"Failed to create PDF: {e}"}

    def create_docx(self, file_path: str, title: str,
                    sections: List[Dict[str, Any]], theme: str = "professional") -> Dict[str, Any]:
        if not HAS_DOCX:
            return {"status": "error", "message": "python-docx not installed"}

        THEMES = {
            "professional": {"title_color": DocxRGB(0x0D,0x1B,0x2A), "heading_color": DocxRGB(0x00,0x72,0xC6), "body_color": DocxRGB(0x33,0x33,0x33), "accent": DocxRGB(0x00,0xB4,0xD8)},
            "modern":       {"title_color": DocxRGB(0x1A,0x1A,0x2E), "heading_color": DocxRGB(0x45,0x52,0xE1), "body_color": DocxRGB(0x22,0x22,0x22), "accent": DocxRGB(0x45,0x52,0xE1)},
            "corporate":    {"title_color": DocxRGB(0x1B,0x26,0x3B), "heading_color": DocxRGB(0xE8,0x63,0x1A), "body_color": DocxRGB(0x22,0x22,0x22), "accent": DocxRGB(0xE8,0x63,0x1A)},
        }
        C = THEMES.get(theme, THEMES["professional"])
        doc = DocxDocument()
        for sec in doc.sections:
            sec.top_margin = sec.bottom_margin = DocxInches(1)
            sec.left_margin = sec.right_margin = DocxInches(1.2)

        # Title
        tp = doc.add_paragraph()
        tr = tp.add_run(title)
        tr.font.size = DocxPt(28); tr.font.bold = True
        tr.font.color.rgb = C["title_color"]; tr.font.name = "Calibri Light"
        tp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        tp.paragraph_format.space_after = DocxPt(4)

        dp = doc.add_paragraph()
        dr = dp.add_run(datetime.datetime.now().strftime("%B %d, %Y"))
        dr.font.size = DocxPt(11); dr.font.color.rgb = C["accent"]; dr.font.name = "Calibri"
        dp.paragraph_format.space_after = DocxPt(16)
        doc.add_paragraph()

        for sec in sections:
            heading_text = sec.get("heading", "")
            if heading_text:
                h = doc.add_paragraph()
                hr = h.add_run(heading_text)
                hr.font.size = DocxPt(16); hr.font.bold = True
                hr.font.color.rgb = C["heading_color"]; hr.font.name = "Calibri"
                h.paragraph_format.space_before = DocxPt(14)
                h.paragraph_format.space_after  = DocxPt(4)
            for para_text in sec.get("paragraphs", []):
                p = doc.add_paragraph()
                r = p.add_run(para_text)
                r.font.size = DocxPt(11); r.font.color.rgb = C["body_color"]; r.font.name = "Calibri"
                p.paragraph_format.space_after = DocxPt(6)
            for bullet in sec.get("bullets", []):
                bp = doc.add_paragraph(style="List Bullet")
                br = bp.add_run(bullet)
                br.font.size = DocxPt(11); br.font.color.rgb = C["body_color"]; br.font.name = "Calibri"

        Path(file_path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(file_path)
        self.active_file = str(Path(file_path).resolve())
        return {"status": "success", "action": "create_docx", "file": file_path,
                "title": title, "section_count": len(sections), "theme": theme}

    def create_xlsx(self, file_path: str, title: str,
                    sheets: List[Dict[str, Any]], theme: str = "dark") -> Dict[str, Any]:
        if not HAS_OPENPYXL:
            return {"status": "error", "message": "openpyxl not installed"}

        THEMES = {
            "dark":      {"hf": "0D1B2A", "hc": "00B4D8", "tc": "00B4D8", "af": "1A2A3A", "bc": "00B4D8"},
            "light":     {"hf": "4552E1", "hc": "FFFFFF", "tc": "1A1A2E", "af": "EEF0FF", "bc": "4552E1"},
            "corporate": {"hf": "1B263B", "hc": "E8631A", "tc": "E8631A", "af": "EEF2F7", "bc": "E8631A"},
        }
        C = THEMES.get(theme, THEMES["dark"])
        thin = Side(style="thin", color=C["bc"])
        bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)
        wb   = Workbook(); wb.remove(wb.active)

        for sdef in sheets:
            ws      = wb.create_sheet(title=sdef.get("name", "Sheet")[:31])
            headers = sdef.get("headers", [])
            rows    = sdef.get("rows", [])
            ncols   = max(len(headers), 1)

            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
            tc = ws.cell(row=1, column=1, value=title)
            tc.font = Font(name="Calibri", size=14, bold=True, color=C["tc"])
            tc.fill = PatternFill("solid", fgColor="111111" if theme == "dark" else "F8F9FA")
            tc.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 28

            for ci, hdr in enumerate(headers, 1):
                c = ws.cell(row=2, column=ci, value=hdr)
                c.font   = Font(name="Calibri", bold=True, color=C["hc"])
                c.fill   = PatternFill("solid", fgColor=C["hf"])
                c.border = bdr
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.column_dimensions[get_column_letter(ci)].width = max(15, len(str(hdr)) + 4)
            if headers: ws.row_dimensions[2].height = 20

            for ri, row in enumerate(rows, 3):
                for ci, val in enumerate(row, 1):
                    c = ws.cell(row=ri, column=ci, value=val)
                    c.font = Font(name="Calibri", size=10); c.border = bdr
                    if ri % 2 == 0: c.fill = PatternFill("solid", fgColor=C["af"])
                    c.alignment = Alignment(wrap_text=True, vertical="center")
                ws.row_dimensions[ri].height = 16

            summary = sdef.get("summary", "")
            if summary:
                sr = len(rows) + 3
                ws.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=ncols)
                sc = ws.cell(row=sr, column=1, value=f"Summary: {summary}")
                sc.font = Font(name="Calibri", italic=True, color=C["tc"])
            ws.freeze_panes = "A3"

        Path(file_path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(file_path)
        self.active_file = str(Path(file_path).resolve())
        return {"status": "success", "action": "create_xlsx", "file": file_path,
                "title": title, "sheet_count": len(sheets), "theme": theme}

    # ── Edit ─────────────────────────────────────────────────────────────

    def edit_docx(self, file_path: str, action: str,
                  heading: Optional[str] = None,
                  paragraphs: Optional[List[str]] = None,
                  bullets: Optional[List[str]] = None,
                  find_text: Optional[str] = None,
                  replace_text_val: Optional[str] = None,
                  table_headers: Optional[List[str]] = None,
                  table_rows: Optional[List[List[str]]] = None) -> Dict[str, Any]:
        if not HAS_DOCX: return {"status": "error", "message": "python-docx not installed"}
        doc = DocxDocument(file_path)
        result = {"status": "success", "action": action, "file": file_path}

        if action == "append_section":
            if heading: doc.add_heading(heading, level=2)
            for p in (paragraphs or []): doc.add_paragraph(p)
            for b in (bullets or []): doc.add_paragraph(b, style="List Bullet")

        elif action == "replace_text":
            count = 0; ft = find_text or ""; rt = replace_text_val or ""
            for para in doc.paragraphs:
                for run in para.runs:
                    if ft in run.text: run.text = run.text.replace(ft, rt); count += 1
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            for run in para.runs:
                                if ft in run.text: run.text = run.text.replace(ft, rt); count += 1
            result["occurrences_replaced"] = count

        elif action == "add_table":
            if table_headers and table_rows is not None:
                tbl = doc.add_table(rows=1 + len(table_rows), cols=len(table_headers))
                tbl.style = "Table Grid"
                for i, h in enumerate(table_headers): tbl.rows[0].cells[i].text = h
                for ri, row in enumerate(table_rows, 1):
                    for ci, val in enumerate(row): tbl.rows[ri].cells[ci].text = str(val)

        doc.save(file_path)
        return result

    def edit_xlsx(self, file_path: str, action: str,
                  sheet_name: Optional[str] = None,
                  rows: Optional[List[List[Any]]] = None,
                  find_value: Optional[str] = None,
                  replace_value: Optional[str] = None,
                  new_sheet_name: Optional[str] = None,
                  headers: Optional[List[str]] = None) -> Dict[str, Any]:
        if not HAS_OPENPYXL: return {"status": "error", "message": "openpyxl not installed"}
        wb = openpyxl.load_workbook(file_path)
        result = {"status": "success", "action": action, "file": file_path}

        if action == "append_rows":
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            for row in (rows or []): ws.append(row)

        elif action == "replace_cell":
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            count = 0
            for row in ws.iter_rows():
                for cell in row:
                    if str(cell.value) == str(find_value): cell.value = replace_value; count += 1
            result["occurrences_replaced"] = count

        elif action == "add_sheet":
            ws = wb.create_sheet(title=(new_sheet_name or "NewSheet")[:31])
            if headers: ws.append(headers)
            for row in (rows or []): ws.append(row)

        wb.save(file_path)
        return result

    def open_file(self, file_path: str) -> bool:
        self.active_file = str(Path(file_path).resolve())
        try:
            os.startfile(os.path.abspath(file_path)); return True
        except Exception as e:
            print(f"[DocAgent] open error: {e}"); return False


# ╔══════════════════════════════════════════════════════════════════════════╗
# ║                  UNIVERSAL RAG INDEX                                    ║
# ╚══════════════════════════════════════════════════════════════════════════╝

class DocRAGIndex:
    def __init__(self, file_path: str, agent: DocAgent):
        self.file_path  = file_path
        self.agent      = agent
        self.chunks: List[Dict[str, Any]] = []
        self.indexed_at = 0.0
        self.doc_type   = DOC_EXTENSIONS.get(Path(file_path).suffix.lower().lstrip('.'), "Unknown")
        self.full_text  = ""
        self._build()

    def _build(self):
        ext = Path(self.file_path).suffix.lower().lstrip('.')
        self.chunks = []
        if ext in ("pptx", "ppt"):   self._index_pptx()
        elif ext in ("docx", "doc"): self._index_docx()
        elif ext in ("xlsx", "xls"): self._index_xlsx()
        elif ext == "pdf":           self._index_pdf()
        self.full_text = "\n\n".join(c["text"] for c in self.chunks)
        self.indexed_at = time.time()

    def _add_chunk(self, cid, ctype, label, text, meta={}):
        tokens = set(re.findall(r"\w+", text.lower()))
        self.chunks.append({"id": cid, "type": ctype, "label": label,
                            "text": text, "tokens": tokens, **meta})

    def _index_pptx(self):
        try:
            prs = pptx.Presentation(self.file_path)
            for i, slide in enumerate(prs.slides, 1):
                title = (slide.shapes.title.text.strip()
                         if slide.shapes.title and slide.shapes.title.has_text_frame else "")
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            t = para.text.strip()
                            if t and t != title: texts.append(t)
                    if shape.has_table:
                        for row in shape.table.rows:
                            cells = [c.text.strip() for c in row.cells if c.text.strip()]
                            if cells: texts.append(" | ".join(cells))
                try:
                    if slide.has_notes_slide:
                        n = slide.notes_slide.notes_text_frame.text.strip()
                        if n: texts.append(f"Notes: {n}")
                except Exception: pass
                full = f"Slide {i}: {title}\n" + "\n".join(texts)
                self._add_chunk(f"slide_{i}", "slide", f"Slide {i}: {title}", full,
                                {"slide_num": i, "title": title, "bullets": texts})
        except Exception as e:
            print(f"[DocRAG] PPTX: {e}")

    def _index_docx(self):
        if not HAS_DOCX: return
        try:
            doc = DocxDocument(self.file_path)
            heading = "Introduction"; buf = []; n = 0

            def flush(h, b):
                nonlocal n
                if b:
                    self._add_chunk(f"para_{n}", "paragraph", h,
                                    f"{h}\n" + "\n".join(b), {"heading": h}); n += 1

            for para in doc.paragraphs:
                t = para.text.strip()
                if not t: continue
                style = para.style.name if para.style else ""
                if "Heading" in style:
                    flush(heading, buf); buf = []; heading = t
                else:
                    buf.append(t)
                    if len(buf) >= 8: flush(heading, buf); buf = []
            flush(heading, buf)

            for i, table in enumerate(doc.tables, 1):
                rows = []
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells: rows.append(" | ".join(cells))
                if rows: self._add_chunk(f"table_{i}", "table", f"Table {i}",
                                         "\n".join(rows), {"table_num": i})
        except Exception as e:
            print(f"[DocRAG] DOCX: {e}")

    def _index_xlsx(self):
        if not HAS_OPENPYXL: return
        try:
            wb = openpyxl.load_workbook(self.file_path, read_only=True, data_only=True)
            for sname in wb.sheetnames:
                ws = wb[sname]; rows = []
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) if c is not None else "" for c in row]
                    if any(c.strip() for c in cells): rows.append(" | ".join(cells))
                    if len(rows) >= 100: break
                if rows: self._add_chunk(f"sheet_{sname}", "sheet", f"Sheet: {sname}",
                                         "\n".join(rows), {"sheet": sname})
            wb.close()
        except Exception as e:
            print(f"[DocRAG] XLSX: {e}")

    def _index_pdf(self):
        if not HAS_PDF: return
        try:
            with pdfplumber.open(self.file_path) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    text = (page.extract_text() or "").strip()
                    if text: self._add_chunk(f"page_{i}", "page", f"Page {i}", text, {"page_num": i})
                    if i >= 50: break
        except Exception as e:
            print(f"[DocRAG] PDF: {e}")

    def search(self, query: str, top_k: int = 5) -> List[Dict[str, Any]]:
        q_tokens = set(re.findall(r"\w+", query.lower()))
        if not q_tokens: return self.chunks[:top_k]
        scored = []
        for chunk in self.chunks:
            overlap = len(q_tokens & chunk["tokens"])
            label_boost = len(q_tokens & set(re.findall(r"\w+", chunk["label"].lower()))) * 3
            score = overlap + label_boost
            if score > 0: scored.append((score, chunk))
        scored.sort(key=lambda x: x[0], reverse=True)
        return [c for _, c in scored[:top_k]]

    def get_summary_text(self) -> str:
        p = Path(self.file_path)
        lines = [f"📄 {self.doc_type}: {p.name} ({len(self.chunks)} sections)\n"]
        for ch in self.chunks[:25]:
            preview = ch["text"][:100].replace("\n", " ")
            lines.append(f"  [{ch['type'].upper()}] {ch['label']}: {preview}…")
        if len(self.chunks) > 25:
            lines.append(f"  … +{len(self.chunks)-25} more sections")
        return "\n".join(lines)


doc_agent = DocAgent()
